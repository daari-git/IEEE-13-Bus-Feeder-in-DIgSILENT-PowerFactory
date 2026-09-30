"""
Build a colour-coded Excel report (Table_II_report.xlsx) from Table_II_results.csv,
which table2_generator.py writes after running in PowerFactory.

  green  : match >= 90 %
  amber  : 75 % <= match < 90 %
  red    : match < 75 %

Run with normal Python (needs openpyxl):  python make_colored_report.py
"""
import csv
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "Table_II_results.csv")
OUT = os.path.join(HERE, "Table_II_report.xlsx")

GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="1F4E78")
BEST = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def colour(pct):
    if pct is None:
        return None
    return GREEN if pct >= 90 else AMBER if pct >= 75 else RED


def num(x):
    try:
        v = float(x)
        return None if v != v else v
    except (TypeError, ValueError):
        return None


def mean(vals):
    vals = [v for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


def header(ws, row, titles):
    for c, t in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def put(ws, row, col, value, fmt=None, fill=None, bold=False):
    cell = ws.cell(row=row, column=col, value=value)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    cell.font = Font(bold=bold)
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center")
    return cell


# ---- read ------------------------------------------------------------------------------
studies = OrderedDict()
with open(SRC, newline="") as fh:
    for r in csv.DictReader(fh):
        studies.setdefault(r["Study"], []).append(r)

scores = []
for study, rows in studies.items():
    a_in = mean([num(r["Inom_match_%"]) for r in rows])
    a_mn = mean([num(r["Ifmin_match_%"]) for r in rows])
    a_mx = mean([num(r["Ifmax_match_%"]) for r in rows])
    scores.append((study, a_in, a_mn, a_mx, mean([a_in, a_mn, a_mx])))
scores.sort(key=lambda s: -(s[4] or 0))
best = scores[0][0]

wb = Workbook()

# ---- sheet 1: Table II of the best study -------------------------------------------------
ws = wb.active
ws.title = "Table II (best match)"
ws["A1"] = "TABLE II - Maximum / minimum / rated branch current, IEEE 13-node feeder"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Best-matching case: " + best
ws["A2"].font = Font(italic=True)
ws["A3"] = "Match % = 100 - |mine - paper| / paper x 100    green >= 90 %   amber 75-90 %   red < 75 %"
header(ws, 5, ["From", "To", "I_nom (A)", "Paper (A)", "Match",
               "I_f,min (kA)", "Paper (kA)", "Match",
               "I_f,max (kA)", "Paper (kA)", "Match", "Compared value"])
row = 6
for r in studies[best]:
    m_in, m_mn, m_mx = num(r["Inom_match_%"]), num(r["Ifmin_match_%"]), num(r["Ifmax_match_%"])
    put(ws, row, 1, r["From"])
    put(ws, row, 2, r["To"])
    put(ws, row, 3, num(r["Inom_A"]), "0.0")
    put(ws, row, 4, num(r["Inom_paper_A"]), "0.0")
    put(ws, row, 5, m_in / 100 if m_in is not None else "n/a", "0%", colour(m_in))
    put(ws, row, 6, num(r["Ifmin_kA"]), "0.00")
    put(ws, row, 7, num(r["Ifmin_paper_kA"]), "0.00")
    put(ws, row, 8, m_mn / 100 if m_mn is not None else "n/a", "0%", colour(m_mn))
    put(ws, row, 9, num(r["Ifmax_kA"]), "0.00")
    put(ws, row, 10, num(r["Ifmax_paper_kA"]), "0.00")
    put(ws, row, 11, m_mx / 100 if m_mx is not None else "n/a", "0%", colour(m_mx))
    put(ws, row, 12, "bus fault current" if r["Compared_value"] == "bus" else "branch current")
    row += 1
s = next(x for x in scores if x[0] == best)
put(ws, row, 1, "AVERAGE", bold=True)
ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
put(ws, row, 5, s[1] / 100, "0%", colour(s[1]), bold=True)
put(ws, row, 8, s[2] / 100, "0%", colour(s[2]), bold=True)
put(ws, row, 11, s[3] / 100, "0%", colour(s[3]), bold=True)
for c, w in enumerate([10, 7, 11, 11, 9, 13, 11, 9, 13, 11, 9, 18], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A6"

# ---- sheet 2: ranking of all studies -------------------------------------------------------
ws2 = wb.create_sheet("Method comparison")
ws2["A1"] = "Average match with the paper for every short-circuit case (best first)"
ws2["A1"].font = Font(bold=True, size=13)
header(ws2, 3, ["Rank", "Short-circuit case", "I_nom", "I_f,min", "I_f,max", "Overall"])
for i, (study, a_in, a_mn, a_mx, a_all) in enumerate(scores, 1):
    r = 3 + i
    put(ws2, r, 1, i, fill=BEST if i == 1 else None, bold=(i == 1))
    c = put(ws2, r, 2, study, fill=BEST if i == 1 else None, bold=(i == 1))
    c.alignment = Alignment(horizontal="left")
    for col, v in ((3, a_in), (4, a_mn), (5, a_mx), (6, a_all)):
        put(ws2, r, col, v / 100 if v is not None else "n/a", "0%", colour(v), bold=(i == 1))
ws2.column_dimensions["A"].width = 6
ws2.column_dimensions["B"].width = 50
for col in "CDEF":
    ws2.column_dimensions[col].width = 10

# ---- sheet 3: every row of every study -------------------------------------------------------
ws3 = wb.create_sheet("All results")
cols = list(next(iter(studies.values()))[0].keys())
header(ws3, 1, cols)
r = 2
for rows in studies.values():
    for rec in rows:
        for c, k in enumerate(cols, 1):
            v = num(rec[k]) if k not in ("Study", "From", "To", "Compared_value") else rec[k]
            fill = colour(v) if k.endswith("match_%") else None
            put(ws3, r, c, v if v is not None else "n/a", "0.0" if k.endswith("%") else None,
                fill)
        r += 1
ws3.column_dimensions["A"].width = 44
ws3.freeze_panes = "B2"

wb.save(OUT)
print("Best match :", best)
print("Averages   : Inom %.0f%% | If,min %.0f%% | If,max %.0f%% | overall %.0f%%" % s[1:])
print("Saved      :", OUT)
